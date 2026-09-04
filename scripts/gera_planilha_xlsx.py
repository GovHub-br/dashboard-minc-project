#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera a planilha formatada dos eixos a partir do CSV versionado.

O CSV é a fonte; este script só o veste. Assim planilha e painel nunca
divergem — os dois saem do mesmo arquivo.

O .xlsx resultante, ao ser enviado ao Google Drive, é convertido em planilha do
Sheets preservando cores, congelamento de cabeçalho e listas suspensas.

Uso:
    python3 scripts/gera_planilha_xlsx.py

Requer openpyxl (só este script; o resto do projeto usa apenas a biblioteca
padrão).
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parent.parent
ORIGEM = RAIZ / "planilha" / "modelo-eixos-pnc.csv"
DESTINO = RAIZ / "planilha" / "Painel dos 8 Eixos do PNC.xlsx"

ROXO = "7A34F3"
CINZA_TEXTO = "2D3748"
BRANCO = "FFFFFF"

# Uma cor por eixo. O tom forte marca a faixa do eixo; o claro pinta as linhas.
# As cores não carregam significado sozinhas — o número do eixo está escrito em
# toda linha, para quem não distingue matizes.
CORES_EIXO = {
    1: ("7A34F3", "EDE4FE"),
    2: ("C026D3", "FAE8FF"),
    3: ("E11D48", "FFE4E9"),
    4: ("F97316", "FFEDD5"),
    5: ("CA8A04", "FEF6DC"),
    6: ("10B981", "DCFCE7"),
    7: ("0D9488", "D3F3EF"),
    8: ("2563EB", "DBEAFE"),
}

COLUNAS = [
    ("Prioridade", 12),
    ("Eixo", 34),
    ("Processo", 20),
    ("Atividade", 18),
    ("Cod_Task", 11),
    ("Tarefa", 62),
    ("Depende de", 20),
    ("Responsável", 15),
    ("Prazo", 17),
    ("Status", 15),
]

PRIORIDADES = ["Alta", "Média", "Baixa"]
STATUS = ["Não iniciado", "Em andamento", "Concluído", "Bloqueado", "Em risco"]

BORDA = Side(style="thin", color="D8D8E0")
GRADE = Border(left=BORDA, right=BORDA, top=BORDA, bottom=BORDA)


def numero_do_eixo(texto):
    texto = (texto or "").strip()
    return int(texto[0]) if texto[:1].isdigit() else 0


def le_csv():
    with open(ORIGEM, encoding="utf-8", newline="") as f:
        linhas = list(csv.reader(f))
    return linhas[0], linhas[1:]


def monta_demandas(wb, cabecalho, dados):
    ws = wb.active
    ws.title = "Demandas"

    for i, (nome, largura) in enumerate(COLUNAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
        c = ws.cell(row=1, column=i, value=nome)
        c.font = Font(bold=True, color=BRANCO, size=11)
        c.fill = PatternFill("solid", fgColor=ROXO)
        c.alignment = Alignment(vertical="center", horizontal="center")
        c.border = GRADE
    ws.row_dimensions[1].height = 30

    for n, linha in enumerate(dados, start=2):
        eixo = numero_do_eixo(linha[1])
        forte, claro = CORES_EIXO.get(eixo, ("999999", "F2F2F2"))

        for i, valor in enumerate(linha, start=1):
            c = ws.cell(row=n, column=i, value=valor)
            c.fill = PatternFill("solid", fgColor=claro)
            c.border = GRADE
            c.alignment = Alignment(vertical="top", wrap_text=(i == 6))
            c.font = Font(size=10, color=CINZA_TEXTO)

        # Coluna do eixo: cor forte, para a faixa ser localizável de relance.
        ce = ws.cell(row=n, column=2)
        ce.font = Font(size=10, bold=True, color=BRANCO)
        ce.fill = PatternFill("solid", fgColor=forte)
        ce.alignment = Alignment(vertical="top", wrap_text=True)

        # Cod_Task como texto: sem isso o Sheets lê "E1.1.1" tudo bem, mas um
        # código sem prefixo digitado depois viraria data.
        cc = ws.cell(row=n, column=5)
        cc.font = Font(size=10, bold=True, color=forte)
        cc.number_format = "@"

        ws.cell(row=n, column=7).number_format = "@"
        ws.row_dimensions[n].height = 30

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{len(dados) + 1}"

    ultima = len(dados) + 400  # espaço para novas linhas herdarem as listas

    dv_prio = DataValidation(
        type="list", formula1=f'"{",".join(PRIORIDADES)}"', allow_blank=True,
        showDropDown=False, promptTitle="Prioridade",
        prompt="Alta, Média ou Baixa. Em branco vira Média.")
    ws.add_data_validation(dv_prio)
    dv_prio.add(f"A2:A{ultima}")

    dv_status = DataValidation(
        type="list", formula1=f'"{",".join(STATUS)}"', allow_blank=True,
        showDropDown=False, promptTitle="Status",
        prompt="Use Bloqueado só para bloqueio que a rede não enxerga. "
               "Travada por dependência o painel calcula sozinho.")
    ws.add_data_validation(dv_status)
    dv_status.add(f"J2:J{ultima}")

    dv_cod = DataValidation(
        type="custom", formula1='TRUE', allow_blank=True,
        promptTitle="Código da demanda",
        prompt="Formato E<eixo>.<bloco>.<n>, por exemplo E1.2.3. "
               "A letra na frente é obrigatória: sem ela a planilha converte "
               "o código em data e a rede de dependências se perde.")
    ws.add_data_validation(dv_cod)
    dv_cod.add(f"E2:E{ultima}")

    dv_prazo = DataValidation(
        type="custom", formula1='TRUE', allow_blank=True,
        promptTitle="Prazo",
        prompt="DD/MM/AAAA. Enquanto não houver data acordada, escreva "
               "'não definido ainda' — o painel trata como sem prazo, nunca "
               "como atraso.")
    ws.add_data_validation(dv_prazo)
    dv_prazo.add(f"I2:I{ultima}")

    return ws


def monta_guia(wb, dados):
    ws = wb.create_sheet("Como preencher")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 96

    def titulo(texto):
        ws.append([])
        n = ws.max_row + 1
        c = ws.cell(row=n, column=1, value=texto)
        c.font = Font(bold=True, size=12, color=ROXO)
        ws.merge_cells(start_row=n, start_column=1, end_row=n, end_column=2)

    def par(rotulo, texto):
        n = ws.max_row + 1
        a = ws.cell(row=n, column=1, value=rotulo)
        a.font = Font(bold=True, size=10, color=CINZA_TEXTO)
        a.alignment = Alignment(vertical="top")
        b = ws.cell(row=n, column=2, value=texto)
        b.font = Font(size=10, color=CINZA_TEXTO)
        b.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[n].height = max(16, 14 * (len(texto) // 90 + 1))

    c = ws.cell(row=1, column=1, value="Painel dos 8 Eixos do PNC — como preencher")
    c.font = Font(bold=True, size=15, color=ROXO)
    ws.merge_cells("A1:B1")

    titulo("O essencial")
    par("A planilha é a fonte",
        "O que estiver aqui aparece no painel; o que não estiver, não existe "
        "para o painel. A sincronização roda todo dia às 6h.")
    par("Uma demanda por linha",
        "Não agrupe duas coisas numa linha só: o painel não consegue separar "
        "o que está pronto do que não está.")

    titulo("As colunas que não podem errar")
    par("Cod_Task",
        "Formato E<eixo>.<bloco>.<n> — E1.2.3. Precisa ser único na planilha "
        "inteira, e o número do eixo tem de bater com a coluna Eixo. É por ele "
        "que as dependências se ligam.")
    par("A letra na frente",
        "Obrigatória. Sem ela o Google Sheets converte 1.1.1 em 01/01/2001, "
        "sem avisar, e a rede de dependências se perde em silêncio.")
    par("Depende de",
        "Os códigos que precisam terminar antes, separados por ponto e "
        "vírgula: E1.1.1; E1.2.3. Deixe vazio quando a demanda pode começar "
        "a qualquer momento.")
    par("Prazo",
        "DD/MM/AAAA. Sem data acordada, escreva 'não definido ainda' — o "
        "painel mostra como sem prazo e nunca como atraso. Prazo em branco e "
        "prazo vencido são coisas diferentes.")
    par("Status",
        "Não iniciado, Em andamento, Concluído, Bloqueado ou Em risco. Use "
        "Bloqueado só para bloqueio que a rede não enxerga — falta de decisão, "
        "de orçamento. Travada por dependência o painel calcula sozinho.")

    titulo("O que o painel calcula sozinho")
    par("Não preencha",
        "Percentual de conclusão, progresso por eixo, quais demandas estão "
        "travadas, quais estão prontas para iniciar, qual é o gargalo, a "
        "cadeia mais longa e a evolução no tempo.")

    titulo("Cores dos eixos")
    for numero in sorted(CORES_EIXO):
        nome = next(
            (l[1] for l in dados if numero_do_eixo(l[1]) == numero), f"Eixo {numero}"
        )
        forte, claro = CORES_EIXO[numero]
        n = ws.max_row + 1
        a = ws.cell(row=n, column=1, value=f"Eixo {numero}")
        a.font = Font(bold=True, size=10, color=BRANCO)
        a.fill = PatternFill("solid", fgColor=forte)
        a.alignment = Alignment(horizontal="center", vertical="center")
        b = ws.cell(row=n, column=2, value=nome)
        b.fill = PatternFill("solid", fgColor=claro)
        b.font = Font(size=10, color=CINZA_TEXTO)

    titulo("Publicar")
    par("Uma vez só",
        "Arquivo → Compartilhar → Publicar na web, formato CSV. Depois "
        "cadastre o ID da planilha no repositório, em Settings → Secrets and "
        "variables → Actions → Variables, com o nome PLANILHA_ID.")
    return ws


def main():
    cabecalho, dados = le_csv()
    wb = Workbook()
    monta_demandas(wb, cabecalho, dados)
    monta_guia(wb, dados)
    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DESTINO)

    eixos = sorted({numero_do_eixo(l[1]) for l in dados})
    print(f"✓ {DESTINO.name} — {len(dados)} demandas, eixos {eixos}.")


if __name__ == "__main__":
    main()
